"""Genera el reporte final editable (DOCX) a partir de los artefactos medidos."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "04_final" / "reporte_final.docx"
BLUE = "174A7E"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None, font_size=8.3):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[c_idx])
            set_cell_fill(cell, BLUE if r_idx == 0 else (LIGHT_BLUE if r_idx % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if c_idx > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)


def add_image(doc, rel, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(ROOT / rel), width=Inches(width))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def main():
    results = pd.read_csv(ROOT / "03_evaluacion" / "resultados.csv")
    eff = json.loads((ROOT / "02_modelo" / "config_efectiva.json").read_text(encoding="utf-8"))
    cfg = json.loads((ROOT / "01_datos" / "config.json").read_text(encoding="utf-8"))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.68)
    sec.right_margin = Inches(0.68)
    sec.header_distance = Inches(0.25)
    sec.footer_distance = Inches(0.25)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.06
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(23)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].paragraph_format.space_after = Pt(8)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].paragraph_format.space_before = Pt(6)
        styles[name].paragraph_format.space_after = Pt(3)
        styles[name].paragraph_format.keep_with_next = True
    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9.2)
    styles["List Bullet"].paragraph_format.space_after = Pt(2)

    for section in doc.sections:
        h = section.header.paragraphs[0]
        h.text = "K-SVD para reconstrucción y eliminación de ruido"
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h.runs[0].font.name = "Aptos"
        h.runs[0].font.size = Pt(7.5)
        h.runs[0].font.color.rgb = RGBColor(90, 90, 90)
        add_page_number(section.footer.paragraphs[0])

    # Página 1
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Reconstrucción y eliminación de ruido en imágenes con K-SVD")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    r = subtitle.add_run("Reporte final de implementación y auditoría técnica")
    r.font.name = "Aptos"
    r.font.size = Pt(12)
    r.font.bold = True
    paragraph(doc, "Proyecto académico de Matemáticas Computacionales · Integración final: Sarah Vásquez M. · 3 de octubre de 2026")

    doc.add_heading("Resumen", level=1)
    paragraph(doc, "Se implementó una adaptación didáctica de K-SVD para aprender un diccionario sobre parches 8×8 de camera y reconstruir coins, una imagen no usada en entrenamiento. El pipeline usa OMP de scikit-learn, actualizaciones de átomos con numpy.linalg.svd y una única implementación modular. En la imagen ruidosa con σ=20/255, el mejor caso fue T0=4 con 26.12 dB, una mejora de 4.03 dB frente a la entrada ruidosa. La auditoría local ejecutó los cuatro bloques, 32 pruebas y los cuatro notebooks sin errores.")

    doc.add_heading("Objetivo y alcance", level=1)
    paragraph(doc, "El objetivo es estudiar cómo la dispersión controla el compromiso entre representación y ajuste de ruido. No se busca reproducir todos los detalles del estimador del artículo ni competir con métodos actuales. La evidencia se limita a una imagen de prueba, una realización de ruido y tres valores fijados de T0.")

    doc.add_heading("Diseño experimental", level=1)
    table = doc.add_table(rows=1, cols=3)
    for i, v in enumerate(("Componente", "Decisión", "Control")):
        table.rows[0].cells[i].text = v
    rows = [
        ("Datos", "camera 256×256 / coins 128×128", "Prueba separada del entrenamiento"),
        ("Parches", "8×8, paso 4, centrados", "Señales por columnas, float64"),
        ("Modelo", "K=128, 8 iteraciones, T=4", "CPU, semillas 42/43/44"),
        ("Evaluación", "T0=2, 4, 8; σ=20/255", "MSE, PSNR, actividad y tiempo"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    style_table(table, widths=[1.25, 2.7, 2.95], font_size=8.5)
    paragraph(doc, "Contribución de Sarah. El bloque de datos se conservó como fuente de verdad y se reforzó con validaciones de finitud, tipos, posiciones, cobertura y parámetros inválidos. Los arrays coinciden exactamente con el paquete privado entregado en Kaggle.", "Contribución de Sarah.")

    doc.add_page_break()

    # Página 2
    doc.add_heading("1 Datos y representación por parches", level=1)
    paragraph(doc, "Las imágenes de scikit-image se convierten con img_as_float64 antes del redimensionamiento. De camera se extraen 3969 parches; los 1500 parches de entrenamiento se eligen sin reemplazo con RNG 42. Coins produce 961 parches y permanece fuera del aprendizaje.")
    add_image(doc, "01_datos/figuras/imagenes_camera_coins.png", 6.15)
    add_caption(doc, "Figura 1. Imágenes de entrenamiento y prueba después del redimensionamiento.")
    paragraph(doc, "Cada parche xᵢ se centra como zᵢ=xᵢ−μᵢ1 y se almacena como una columna de 64 elementos en orden C. La media μᵢ se conserva para el reensamblado. El control extraer → restituir media → ensamblar recuperó ambas imágenes con error máximo 1.11×10⁻¹⁶.")

    doc.add_heading("2 Codificación dispersa y K-SVD", level=1)
    paragraph(doc, "Para un diccionario D∈R⁶⁴ˣ¹²⁸, OMP obtiene coeficientes aᵢ con a lo más T entradas activas. En cada iteración K-SVD se alternan la codificación y la actualización de cada átomo. Para el átomo j se restringe el residual a las señales que lo usan, se calcula Eⱼ=UΣVᵀ y se asignan dⱼ=u₁ y aⱼ,ω=σ₁v₁ᵀ. Esta SVD entrega la mejor aproximación de rango uno sobre el soporte fijo.")
    bullet(doc, "OMP: sklearn.linear_model.orthogonal_mp; las paradas tempranas documentadas significan que T es un máximo.")
    bullet(doc, "SVD: numpy.linalg.svd sobre el residual restringido; no se usa DictionaryLearning.")
    bullet(doc, "Los parches y reconstrucciones permanecen en float64 y no se aplica clipping.")

    doc.add_page_break()

    # Página 3
    doc.add_heading("3 Entrenamiento del diccionario", level=1)
    paragraph(doc, f"La configuración completa se ejecutó sin contingencia: K={eff['n_atoms']}, {eff['n_train_patches']} parches, {eff['n_iter']} iteraciones y T={eff['train_sparsity']}. El MSE de entrenamiento bajó de {eff['train_mse_init_dictionary']:.3e} con el diccionario inicial a {eff['train_mse_final_recoded']:.3e} con el diccionario final recodificado. El tiempo local fue {eff['training_time_s']:.2f} s y depende de la máquina.")
    add_image(doc, "04_final/figuras/diccionarios.png", 6.75)
    add_caption(doc, "Figura 2. Diccionario inicial y diccionario aprendido, 128 átomos de 8×8.")
    add_image(doc, "04_final/figuras/curva_entrenamiento.png", 5.5)
    add_caption(doc, "Figura 3. Error de representación después de cada actualización K-SVD.")
    paragraph(doc, "La curva desciende en las ocho iteraciones medidas. No se interpreta el aspecto visual de cada átomo como una métrica de calidad; la evidencia principal es el error de representación y la evaluación sobre coins.")

    doc.add_page_break()

    # Página 4
    doc.add_heading("4 Resultados de reconstrucción", level=1)
    base = results[results.case_id == "noisy_baseline"].iloc[0]
    noisy = results[(results.method == "ksvd") & (results.input_kind == "noisy")]
    best = noisy.loc[noisy.psnr_db.idxmax()]
    paragraph(doc, f"Todas las métricas usan la imagen coins limpia como referencia y data_range=1.0. La entrada ruidosa obtuvo {base.psnr_db:.2f} dB. El mejor resultado ruidoso fue T0={int(best.max_nonzero)} con {best.psnr_db:.2f} dB y MSE {best.mse:.6f}.")
    table = doc.add_table(rows=1, cols=4)
    for i, v in enumerate(("Entrada", "T0", "MSE", "PSNR dB")):
        table.rows[0].cells[i].text = v
    for _, row in results.iterrows():
        cells = table.add_row().cells
        label = "ruidosa sin procesar" if row.case_id == "noisy_baseline" else ("limpia" if row.input_kind == "clean" else "ruidosa")
        vals = (label, "—" if row.max_nonzero == 0 else str(int(row.max_nonzero)), f"{row.mse:.6f}", f"{row.psnr_db:.2f}")
        for i, v in enumerate(vals):
            cells[i].text = v
    style_table(table, widths=[2.45, 0.8, 1.5, 1.5], font_size=8.2)
    add_image(doc, "04_final/figuras/reconstrucciones.png", 6.75)
    add_caption(doc, "Figura 4. Referencia, entrada ruidosa y reconstrucciones para T0=2, 4 y 8.")
    paragraph(doc, "En la imagen limpia, el PSNR aumenta al permitir más átomos. En la imagen ruidosa el máximo ocurre en T0=4: con T0=2 se pierde detalle y con T0=8 OMP también representa parte del ruido. Esta última explicación es una interpretación consistente con la curva, no una medición causal independiente.")

    doc.add_page_break()

    # Página 5
    doc.add_heading("5 Conclusiones, auditoría y limitaciones", level=1)
    bullet(doc, f"La mejor reconstrucción ruidosa mejora {best.psnr_db - base.psnr_db:.2f} dB sobre la entrada; la ganancia es real para este caso y no se presenta como estado del arte.")
    bullet(doc, "El pipeline conserva los contratos esenciales: columnas float64, diccionario normalizado, máximo de coeficientes activos y ensamblado sin clipping.")
    bullet(doc, "La auditoría corrigió aliases del contrato de evaluación sin retirar los nombres usados por la integración existente.")
    bullet(doc, "Las cuatro matrices del dataset son idénticas, elemento a elemento, a las del paquete aceptado de Sarah.")

    doc.add_heading("Reproducibilidad", level=2)
    paragraph(doc, "La evidencia local incluye python run_all.py, 32 pruebas y ejecución limpia de los cuatro notebooks. El bloque 1 también cuenta con una versión privada Successful en Kaggle CPU; la ejecución completa de los cuatro bloques no se atribuye a Kaggle. Los tiempos cambian entre equipos, mientras que datos, métricas y reconstrucciones se reproducen salvo redondeo de álgebra lineal.")

    doc.add_heading("Limitaciones", level=2)
    paragraph(doc, "Sólo se usa una imagen de entrenamiento, una imagen de prueba, una semilla de modelo y una realización de ruido. No existe conjunto de validación ni se comparan otros algoritmos. El barrido T0∈{2,4,8} fue fijado antes de ejecutar; por tanto, el resultado sirve como prueba de concepto y no como estimación generalizable.")

    doc.add_heading("Referencias", level=2)
    refs = [
        "Aharon, M., Elad, M. y Bruckstein, A. (2006). K-SVD: An Algorithm for Designing Overcomplete Dictionaries for Sparse Representation. IEEE Transactions on Signal Processing, 54(11). https://doi.org/10.1109/TSP.2006.881199",
        "Elad, M. y Aharon, M. (2006). Image Denoising Via Sparse and Redundant Representations Over Learned Dictionaries. IEEE Transactions on Image Processing, 15(12). https://doi.org/10.1109/TIP.2006.881969",
        "scikit-learn. orthogonal_mp. https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.orthogonal_mp.html",
        "NumPy. numpy.linalg.svd. https://numpy.org/doc/stable/reference/generated/numpy.linalg.svd.html",
        "scikit-image. Data and metrics APIs. https://scikit-image.org/docs/stable/api/",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"[{i}] {ref}")

    doc.core_properties.title = "Reconstrucción y eliminación de ruido en imágenes con K-SVD"
    doc.core_properties.subject = "Reporte final de implementación y auditoría técnica"
    doc.core_properties.author = "Sarah Vásquez M."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
